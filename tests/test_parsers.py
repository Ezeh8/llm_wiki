import io

import pytest

try:
    import pymupdf
except ImportError:  # pragma: no cover
    import fitz as pymupdf

from docx import Document

from app.parsers import (
    PARSE_ERROR_MESSAGE,
    BlockType,
    FileParseError,
    parse_docx,
    parse_document,
    parse_md,
    parse_pdf,
    parse_txt,
)


# --- fixtures -----------------------------------------------------------------


def _make_pdf(text: str = "Hello world paragraph.", *, encrypt: bool = False) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    if encrypt:
        return doc.tobytes(
            encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="o", user_pw="u"
        )
    return doc.tobytes()


def _make_docx() -> bytes:
    document = Document()
    document.add_heading("Access Policy", level=1)
    document.add_paragraph("Badges are reissued at the front desk.")
    document.add_heading("Escalation", level=2)
    document.add_paragraph("Contact security for lost cards.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


_MARKDOWN = b"""# Title

Intro paragraph one.

## Section

Body text here.

```python
# not a heading, inside a fence
x = 1
```
"""

_TXT = b"First paragraph line one.\nline two.\n\nSecond paragraph."


# --- happy paths --------------------------------------------------------------


def test_parse_txt_detects_paragraphs():
    doc = parse_txt(_TXT)
    assert doc.file_type == "txt"
    assert [b.type for b in doc.blocks] == [BlockType.PARAGRAPH, BlockType.PARAGRAPH]
    assert doc.blocks[0].text == "First paragraph line one.\nline two."
    assert doc.blocks[1].text == "Second paragraph."


def test_parse_md_extracts_headings_levels_and_ignores_fenced_hash():
    doc = parse_md(_MARKDOWN)
    kinds = [(b.type, b.level, b.text) for b in doc.blocks]
    assert (BlockType.HEADING, 1, "Title") in kinds
    assert (BlockType.HEADING, 2, "Section") in kinds
    # The '#' inside the code fence must remain part of a paragraph block, not a heading.
    headings = [b for b in doc.blocks if b.type == BlockType.HEADING]
    assert [h.text for h in headings] == ["Title", "Section"]
    fenced = [b for b in doc.blocks if "x = 1" in b.text]
    assert len(fenced) == 1 and fenced[0].type == BlockType.PARAGRAPH


def test_parse_docx_maps_heading_styles_to_levels():
    doc = parse_docx(_make_docx())
    assert doc.file_type == "docx"
    heading = next(b for b in doc.blocks if b.type == BlockType.HEADING)
    assert heading.text == "Access Policy" and heading.level == 1
    assert any(
        b.type == BlockType.HEADING and b.level == 2 and b.text == "Escalation"
        for b in doc.blocks
    )
    assert any("front desk" in b.text for b in doc.blocks)


def test_parse_pdf_extracts_paragraphs_with_page_numbers():
    doc = parse_pdf(_make_pdf("Reset your badge at the front desk."))
    assert doc.file_type == "pdf"
    assert doc.blocks
    assert doc.blocks[0].page == 1
    assert "front desk" in doc.blocks[0].text


def test_registry_dispatch_and_alias():
    assert parse_document("md", _MARKDOWN).file_type == "md"
    assert parse_document("markdown", _MARKDOWN).file_type == "md"
    assert parse_document(".TXT", _TXT).file_type == "txt"


def test_registry_unknown_type_raises_value_error():
    with pytest.raises(ValueError):
        parse_document("csv", b"a,b,c")


# --- failure paths (typed error, client-safe message) -------------------------


@pytest.mark.parametrize(
    "parser, data",
    [
        (parse_pdf, b"%PDF-1.4 not really a pdf"),
        (parse_docx, b"PK-but-not-a-real-docx"),
        (parse_txt, b"\xff\xfe\x00\x81 invalid utf-8"),
        (parse_md, b"\xff\xfe\x00\x81 invalid utf-8"),
    ],
)
def test_corrupted_file_raises_typed_error(parser, data):
    with pytest.raises(FileParseError) as exc:
        parser(data)
    assert str(exc.value) == PARSE_ERROR_MESSAGE


def test_password_protected_pdf_raises_typed_error():
    with pytest.raises(FileParseError) as exc:
        parse_pdf(_make_pdf(encrypt=True))
    assert str(exc.value) == PARSE_ERROR_MESSAGE
