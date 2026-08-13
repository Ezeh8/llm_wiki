import io
import re

from docx import Document

from app.parsers.base import Block, BlockType, FileParseError, ParsedDocument

_HEADING_RE = re.compile(r"^Heading (\d+)$")


def _heading_level(style_name: str) -> int | None:
    if style_name == "Title":
        return 1
    match = _HEADING_RE.match(style_name)
    return int(match.group(1)) if match else None


def parse_docx(data: bytes) -> ParsedDocument:
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise FileParseError() from exc

    try:
        blocks: list[Block] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = getattr(paragraph.style, "name", "") or ""
            level = _heading_level(style_name)
            if level is not None:
                blocks.append(Block(type=BlockType.HEADING, text=text, level=level))
            else:
                blocks.append(Block(type=BlockType.PARAGRAPH, text=text))
    except Exception as exc:
        raise FileParseError() from exc

    return ParsedDocument(file_type="docx", blocks=blocks)
